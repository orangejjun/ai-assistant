from pathlib import Path
from typing import Callable


def parse(file_path: str) -> str:
    """
    파일 형식을 자동 감지하여 텍스트를 추출한다.

    Args:
        file_path: 파싱할 파일의 경로

    Returns:
        str: 추출된 원본 텍스트
    """
    ext = Path(file_path).suffix.lower()
    _parsers: dict[str, Callable[[str], str]] = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".txt": _parse_txt,
    }
    parser_fn = _parsers.get(ext)
    if parser_fn is None:
        raise ValueError(f"지원하지 않는 파일 형식: {ext}")
    return parser_fn(file_path)


def _parse_pdf(file_path: str) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(file_path)
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(pages)


def _parse_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join(str(cell) if cell is not None else "" for cell in row)
            if line.strip():
                rows.append(line)
    wb.close()
    return "\n".join(rows)


def _parse_txt(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8", errors="replace")


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("사용법: python parsers.py <파일경로>")
        raise SystemExit(1)

    path = sys.argv[1]
    text = parse(path)
    print(f"추출 완료: {len(text)}자")
    print("-" * 40)
    print(text[:500])
