import pytest
from pathlib import Path


@pytest.fixture(autouse=True)
def mock_env(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("OPENAI_API_KEY", "test-openai-key")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-anthropic-key")


@pytest.fixture(scope="session")
def fixtures_dir() -> Path:
    return Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def sample_txt_path(fixtures_dir: Path) -> str:
    return str(fixtures_dir / "sample.txt")


@pytest.fixture(scope="session")
def sample_docx_path(tmp_path_factory: pytest.TempPathFactory) -> str:
    from docx import Document

    tmp = tmp_path_factory.mktemp("docs")
    path = tmp / "sample.docx"
    doc = Document()
    doc.add_paragraph("첫 번째 문단입니다. 이 문서는 DOCX 파서 테스트용입니다.")
    doc.add_paragraph("두 번째 문단입니다. 다양한 내용을 포함하고 있습니다.")
    doc.add_paragraph("세 번째 문단입니다. 마지막 내용입니다.")
    doc.save(str(path))
    return str(path)


@pytest.fixture(scope="session")
def sample_xlsx_path(tmp_path_factory: pytest.TempPathFactory) -> str:
    from openpyxl import Workbook

    tmp = tmp_path_factory.mktemp("docs")
    path = tmp / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["이름", "부서", "직급"])
    ws.append(["홍길동", "개발팀", "사원"])
    ws.append(["김철수", "영업팀", "대리"])
    ws.append(["이영희", "인사팀", "과장"])
    wb.save(str(path))
    return str(path)
