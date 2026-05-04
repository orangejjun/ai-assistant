from pathlib import Path
from unittest.mock import patch, MagicMock

import pytest


class TestParseTxt:
    def test_returns_string(self, sample_txt_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_txt_path)

        assert isinstance(result, str)

    def test_contains_expected_content(self, sample_txt_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_txt_path)

        assert "테크솔루션" in result

    def test_nonempty_for_nonempty_file(self, sample_txt_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_txt_path)

        assert len(result) > 0

    def test_handles_utf8_korean(self, tmp_path: Path) -> None:
        from backend.ingest.parsers import parse

        f = tmp_path / "korean.txt"
        f.write_text("가나다라마바사", encoding="utf-8")

        result = parse(str(f))

        assert result == "가나다라마바사"

    def test_empty_file_returns_empty_string(self, tmp_path: Path) -> None:
        from backend.ingest.parsers import parse

        f = tmp_path / "empty.txt"
        f.write_text("")

        result = parse(str(f))

        assert result == ""


class TestParseDocx:
    def test_returns_string(self, sample_docx_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_docx_path)

        assert isinstance(result, str)

    def test_contains_paragraph_text(self, sample_docx_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_docx_path)

        assert "첫 번째 문단" in result
        assert "두 번째 문단" in result

    def test_empty_paragraphs_excluded(self, tmp_path: Path) -> None:
        from docx import Document
        from backend.ingest.parsers import parse

        path = tmp_path / "with_empty.docx"
        doc = Document()
        doc.add_paragraph("내용 있는 문단")
        doc.add_paragraph("")  # 빈 문단
        doc.add_paragraph("   ")  # 공백만 있는 문단
        doc.save(str(path))

        result = parse(str(path))

        assert "내용 있는 문단" in result
        assert result.count("\n") == 0  # 빈 문단은 제외되므로 줄바꿈 없음


class TestParseXlsx:
    def test_returns_string(self, sample_xlsx_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_xlsx_path)

        assert isinstance(result, str)

    def test_contains_cell_values(self, sample_xlsx_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_xlsx_path)

        assert "홍길동" in result
        assert "개발팀" in result

    def test_contains_header_row(self, sample_xlsx_path: str) -> None:
        from backend.ingest.parsers import parse

        result = parse(sample_xlsx_path)

        assert "이름" in result
        assert "부서" in result

    def test_multiple_sheets_parsed(self, tmp_path: Path) -> None:
        from openpyxl import Workbook
        from backend.ingest.parsers import parse

        path = tmp_path / "multi.xlsx"
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1.append(["시트1 데이터"])
        ws2 = wb.create_sheet("Sheet2")
        ws2.append(["시트2 데이터"])
        wb.save(str(path))

        result = parse(str(path))

        assert "시트1 데이터" in result
        assert "시트2 데이터" in result


class TestParsePdf:
    @patch("backend.ingest.parsers._parse_pdf")
    def test_routes_to_pdf_parser(self, mock_parse_pdf: MagicMock, tmp_path: Path) -> None:
        from backend.ingest.parsers import parse

        mock_parse_pdf.return_value = "PDF 내용"
        fake_pdf = tmp_path / "doc.pdf"
        fake_pdf.write_bytes(b"%PDF-1.4 fake")

        result = parse(str(fake_pdf))

        mock_parse_pdf.assert_called_once_with(str(fake_pdf))
        assert result == "PDF 내용"

    def test_parse_pdf_uses_fitz(self, tmp_path: Path) -> None:
        fake_pdf = tmp_path / "doc.pdf"
        fake_pdf.write_bytes(b"%PDF-1.4 fake")

        mock_page = MagicMock()
        mock_page.get_text.return_value = "페이지 텍스트"
        mock_doc = MagicMock()
        mock_doc.__iter__ = MagicMock(return_value=iter([mock_page]))

        with patch("fitz.open", return_value=mock_doc):
            from backend.ingest.parsers import _parse_pdf
            result = _parse_pdf(str(fake_pdf))

        assert "페이지 텍스트" in result


class TestParseUnsupported:
    def test_raises_value_error_for_unknown_extension(self, tmp_path: Path) -> None:
        from backend.ingest.parsers import parse

        f = tmp_path / "data.csv"
        f.write_text("a,b,c")

        with pytest.raises(ValueError, match="지원하지 않는 파일 형식"):
            parse(str(f))

    def test_error_message_contains_extension(self, tmp_path: Path) -> None:
        from backend.ingest.parsers import parse

        f = tmp_path / "log.log"
        f.write_text("log content")

        with pytest.raises(ValueError, match=r"\.log"):
            parse(str(f))
